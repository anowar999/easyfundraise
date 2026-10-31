import io
import json
from datetime import datetime

import openpyxl

from app.services import tracker_import as ti


def make_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    intro = wb.active
    intro.title = "Intro"
    intro["B5"] = "A note about this template"
    ws = wb.create_sheet("Investor Tracker")
    ws.append([])
    ws.append([None, "Wave", "Prio", "Investment firm", "Investor type", "Contact person", "Linked In", "Can lead",
               "Who can intro", "Last meeting date", "Next touch date", "Status (10 days rule)", "Stage",
               "Check size (range)", "Notes (desirability, value add, focus)", "Next steps / TODOs"])
    ws.append([1, 1, "A", "MOMENTOUS", "VC-Partner", "Justin S", None, False, None, datetime(2026, 10, 2),
               datetime(2026, 10, 9), "Active", "5 - Deep Dive / DD", 50000, "Highly interested", "Send materials"])
    ws.append([2, 1, "B", "PLUS9", "VC-Partner", "Raffaele G", "https://linkedin.com/in/ada/", False, "Sam", None, None,
               "Rejection", "4 - Deck requested", None, "Asked for deck", None])
    ws.append([3, 2, "C", "Family Co", "Family Office", "Fam Person", None, True, None, None, None, None,
               "0 - Not Started", None, None, None])
    out = wb.create_sheet("Outreach")
    out.append([])
    out.append([None, None, "Contact person", None, None, None, "Linked In", "Attempts", "Message sent"])
    out.append([None, 1, "Ali C", "ali@example.com", None, "Email", "https://www.linkedin.com/in/ali/", 1, datetime(2026, 10, 12)])
    out.append([None, 2, "Bea D", "bea@example.com", None, "Linkedin", "https://www.linkedin.com/in/bea/", 1, datetime(2026, 10, 12)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_detects_header_and_mapping():
    sheets = {s.name: s for s in ti.parse("t.xlsx", make_xlsx())}
    tr = sheets["Investor Tracker"]
    assert tr.header_row == 1
    fields = set(tr.mapping.values())
    assert {"name", "firm", "investor_type", "priority", "stage", "status", "next_contact", "notes", "tags"} <= fields
    outreach = sheets["Outreach"]
    assert "email" in outreach.mapping.values()  # unlabelled column inferred from values
    assert "contact_method" in outreach.mapping.values()
    assert max(sheets.values(), key=ti.score).name == "Investor Tracker"


def test_value_mapping():
    assert ti.map_stage("0 - Not Started") == 1
    assert ti.map_stage("3 - First Call Arranged") == 4
    assert ti.map_stage("3 - In Contact") == 3
    assert ti.map_stage("4 - Deck requested") == 5
    assert ti.map_stage("5 - Deep Dive / DD") == 6
    assert ti.map_stage("6 - (TS) Offer") == 7
    assert ti.map_stage("7 - Signed") == 8
    assert ti.map_priority("A") == "P1" and ti.map_priority("C") == "P3"
    assert ti.map_type("VC-Partner") == "VC Partner" and ti.map_type("Private Equity") == "PE/Growth"
    assert ti.money(50000) == "$50K"
    assert ti.to_date(45000).year == 2023


def test_import_into_funnel(client):
    data = make_xlsx()
    prev = client.post("/funnel/import/preview", files={"file": ("tracker.xlsx", data)}).json()
    assert prev["suggested_sheet"] == "Investor Tracker"
    sheet = next(s for s in prev["sheets"] if s["name"] == "Investor Tracker")
    mapping = {str(c["index"]): c["field"] for c in sheet["columns"] if c["field"]}
    r = client.post("/funnel/import/commit", params={"sheet": "Investor Tracker", "mapping": json.dumps(mapping)},
                    files={"file": ("tracker.xlsx", data)})
    assert r.status_code == 200, r.text
    res = r.json()
    assert res["created"] == 3 and res["linked"] == 1  # PLUS9 row matches database investor by LinkedIn
    cards = {c["name"]: c for c in client.get("/funnel").json()["items"]}
    j = cards["Justin S"]
    assert j["stage"] == 6 and j["priority"] == "P1" and j["ticket_size"] == "$50K" and "Wave 1" in j["tags"]
    assert j["next_contact_date"] == "2026-10-09"
    assert cards["Raffaele G"]["status"] == "Rejection" and cards["Raffaele G"]["referrer"] == "Sam"
    assert cards["Fam Person"]["investor_type"] == "Family Office" and cards["Fam Person"]["can_lead"] == "Yes"
    notes = client.get(f"/funnel/{j['id']}").json()["events"]
    assert any(e["body"] == "Highly interested" for e in notes)
    # Re-import is idempotent: updates, no duplicates
    r2 = client.post("/funnel/import/commit", params={"sheet": "Investor Tracker", "mapping": json.dumps(mapping)},
                     files={"file": ("tracker.xlsx", data)}).json()
    assert r2["created"] == 0 and r2["updated"] == 3
    assert len(client.get("/funnel").json()["items"]) == 3


def test_import_csv(client):
    csv_bytes = b"Name,Firm,Email,Stage\nCsv Person,Acme VC,csv@acme.vc,In contact\n"
    prev = client.post("/funnel/import/preview", files={"file": ("t.csv", csv_bytes)}).json()
    cols = {c["field"] for c in prev["sheets"][0]["columns"]}
    assert {"name", "firm", "email", "stage"} <= cols
    mapping = {str(c["index"]): c["field"] for c in prev["sheets"][0]["columns"] if c["field"]}
    r = client.post("/funnel/import/commit", params={"sheet": "CSV", "mapping": json.dumps(mapping)},
                    files={"file": ("t.csv", csv_bytes)}).json()
    assert r["created"] == 1
    assert client.get("/funnel").json()["items"][0]["stage"] == 3


def test_rejects_other_files(client):
    r = client.post("/funnel/import/preview", files={"file": ("x.pdf", b"%PDF")})
    assert r.status_code == 415
